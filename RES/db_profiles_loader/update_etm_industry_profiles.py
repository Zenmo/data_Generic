"""Write the ETM industry and ICT electricity curves into db_profiles.xlsx.

Columns
-------
industry_other_e_demand (replaces the existing column)
    Sum of all `industry_*.input (MW)` curves in ETM `electricity_profiles`,
    minus steel (`industry_steel_*`, own column industry_steel_e_demand) and
    minus ICT (own column below).
ict_e_demand (new column, appended after the last column)
    `industry_final_demand_for_other_ict_electricity.input (MW)`: ICT /
    datacenters. Nearly flat: ~5% lower at night, ~4% higher in the day,
    weekends ~3% lower, peak 1.15 x mean. Meant for a separate datacenter asset in
    LUX (see PLAN_profielen_ICT_export.md in the Digital-twin-Drechtsteden repo).
Both are normalised to an annual sum of 1.

Why industry_other_e_demand is replaced
---------------------------------------
The previous column was put in by hand and was wrong in two ways:
  - its weekly dip fell on Monday/Tuesday instead of Saturday/Sunday, in every
    year (the source was two days off before it was laid on the calendar);
  - its shape was not the current ETM curve (correlation 0.28 with the ETM
    industry demand in the same 2023 hours, at best 0.77 when shifted).
house_e_demand_other and building_e_demand_other ARE the ETM curves on the right
calendar (correlation 0.99 / 1.00) and are left alone.

Source
------
Hourly curve `electricity_profiles` of an existing ETM scenario, downloaded with
a plain GET from the public ETM API (no pyetm, no scenario is created).
Default: scenario 1451042, the 2023 present-day scenario of RES region
Drechtsteden (area ES18_drechtsteden). Use --scenario for another region.

The scenario only matters through the subsector mix of industry_other_e_demand:
ETM uses two shapes for industry electricity, one template for food / paper /
other non-specified (peak = 1.99 x mean) and a flat line for chemicals and
metals. The timing is the same for every scenario; the amplitude depends on the
flat share (Drechtsteden 2023: 16% chemicals, so peaks are ~16% lower than the
pure template). The ICT column is a single curve and does not depend on the mix.

Artifact: in ETM the ICT curve is 0 in hour 0 (1 Jan 00:00). That hour is
replaced by the same hour one week later (same weekday).

Calendar
--------
The ETM curves follow the weekdays of 2023 (1 Jan = Sunday; verified: the weekly
dip of industry falls on Saturday/Sunday). Other years are shifted by whole
days, cyclically, so weekdays line up: target day i takes ETM day (i + k) mod 365,
with k chosen so both are the same weekday (2024: k = +1, 2025: k = +3). The
366-day sheet of 2024 gets 366 such days. Public holidays are not moved.

The hour index is used as-is (ETM hour h -> t_h = h), the same convention as the
existing house/building columns (best correlation at shift 0). Quarter-hour
sheets hold the hourly fractions linearly interpolated (sum ~4, like the existing
quarter-hour columns), the last hour held flat, as in
fetch_etm_weather_years.interpolate_to_quarter_hour.

Usage (from this folder):
    python update_etm_industry_profiles.py              # write ../../db_profiles.xlsx
    python update_etm_industry_profiles.py --dry-run    # only fetch + checks
    python update_etm_industry_profiles.py --scenario 1234567

Needs: requests, numpy, pandas, openpyxl.
"""

from __future__ import annotations

import argparse
import io
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import requests

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parent.parent / "db_profiles.xlsx"
DEFAULT_SCENARIO = 1451042
ETM_URL = "https://engine.energytransitionmodel.com/api/v3/scenarios/{id}/curves/electricity_profiles.csv"
ETM_CALENDAR_YEAR = 2023
ICT_COL = "industry_final_demand_for_other_ict_electricity.input (MW)"

# column -> (definition for the Documentation sheet, minimum weekend dip required)
# The weekend check: Saturday and Sunday must be the two lowest days and the
# weekend mean below this fraction of the weekday mean.
COLUMNS = {
    "industry_other_e_demand": (
        "Normalized electricity demand profile for other industry (excl. steel and ICT).", 0.95),
    "ict_e_demand": (
        "Normalized electricity demand profile for ICT / datacenters.", 0.99),
}

# sheet -> (calendar year, number of days, quarter-hour?)
SHEETS = {
    "profiles_2023_h":          (2023, 365, False),
    "profiles_2023":            (2023, 365, True),
    "profiles_2024_h":          (2024, 365, False),
    "profiles_2024":            (2024, 365, True),
    "profiles_2024_366_days_h": (2024, 366, False),
    "profiles_2024_366_days":   (2024, 366, True),
    "profiles_2025_h":          (2025, 365, False),
    "profiles_2025":            (2025, 365, True),
}


def fetch_curves(scenario_id: int) -> dict[str, tuple[np.ndarray, list[str]]]:
    """Hourly curves (MW, 8760 values) per output column, with their source columns."""
    r = requests.get(ETM_URL.format(id=scenario_id), timeout=300)
    r.raise_for_status()
    df = pd.read_csv(io.StringIO(r.text))
    if len(df) != 8760:
        sys.exit(f"Expected 8760 hours, got {len(df)}")
    industry = [c for c in df.columns
                if c.startswith("industry_") and c.endswith(".input (MW)")
                and not c.startswith("industry_steel_") and "_ict_" not in c]
    if not industry or ICT_COL not in df.columns:
        sys.exit("Industry or ICT columns missing in electricity_profiles")
    ict = df[ICT_COL].to_numpy(dtype=float).copy()
    if ict[0] == 0:  # ETM artifact, see module docstring
        ict[0] = ict[7 * 24]
    return {
        "industry_other_e_demand": (df[industry].sum(axis=1).to_numpy(dtype=float), industry),
        "ict_e_demand": (ict, [ICT_COL]),
    }


def day_offset(year: int) -> int:
    """Days to add to a target day to land on the same weekday in the ETM year."""
    return (date(year, 1, 1).weekday() - date(ETM_CALENDAR_YEAR, 1, 1).weekday()) % 7


def for_year(curve: np.ndarray, year: int, days: int) -> np.ndarray:
    """Hourly profile (annual sum 1) for a calendar year, weekdays aligned."""
    k = day_offset(year)
    by_day = curve.reshape(365, 24)
    out = np.concatenate([by_day[(i + k) % 365] for i in range(days)])
    return out / out.sum()


def to_quarter(hourly: np.ndarray) -> np.ndarray:
    t = np.arange(0, len(hourly), 0.25)
    return np.interp(t, np.arange(len(hourly) + 1), np.append(hourly, hourly[-1]))


def weekday_means(hourly: np.ndarray, year: int) -> list[float]:
    """Mean of the daily totals per weekday (Mon..Sun), relative to the overall mean."""
    days = len(hourly) // 24
    daily = hourly.reshape(days, 24).sum(axis=1)
    wd = np.array([(date(year, 1, 1).toordinal() + i) % 7 for i in range(days)])
    # date.toordinal() % 7: 0 = Sunday ... 6 = Saturday -> reorder to Mon..Sun
    mon_first = [1, 2, 3, 4, 5, 6, 0]
    return [float(daily[wd == w].mean() / daily.mean()) for w in mon_first]


def check(column: str, sheet: str, profile: np.ndarray, year: int,
          reference: np.ndarray | None) -> list[str]:
    """Returns a list of problems (empty = OK) and prints the key figures."""
    problems = []
    hourly = profile if len(profile) in (8760, 8784) else profile[::4]
    total = profile.sum() if hourly is profile else profile.sum() / 4
    wm = weekday_means(hourly, year)
    weekend, weekdays = (wm[5] + wm[6]) / 2, sum(wm[:5]) / 5
    line = (f"  {sheet:26s} sum {total:.6f}  Mon..Sun "
            + " ".join(f"{v:.2f}" for v in wm) + f"  weekend/weekday {weekend / weekdays:.3f}")
    if reference is not None:
        corr = np.corrcoef(hourly[:8760], reference)[0, 1]
        line += f"  corr ETM {corr:.4f}"
        if corr < 0.999:
            problems.append(f"{column}/{sheet}: correlation with ETM curve {corr:.4f} < 0.999")
    print(line)
    # Quarter-hour sheets: interpolation shifts a few millionths at the year edges.
    if abs(total - 1) > (1e-6 if hourly is profile else 1e-4):
        problems.append(f"{column}/{sheet}: annual sum {total:.6f} != 1")
    if not weekend < COLUMNS[column][1] * weekdays:
        problems.append(f"{column}/{sheet}: weekend/weekday {weekend / weekdays:.3f} "
                        f"not below {COLUMNS[column][1]}")
    if not (wm[5] < min(wm[:5]) and wm[6] < min(wm[:5])):
        problems.append(f"{column}/{sheet}: Saturday/Sunday are not the two lowest days")
    return problems


def write_column(ws, column: str, values: np.ndarray) -> None:
    """Overwrite the column, or append it after the last header if it is new."""
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if column in header:
        col = header.index(column) + 1
    else:
        col = max(i for i, h in enumerate(header) if h is not None) + 2
        ws.cell(row=1, column=col, value=column)
    n = 0
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        if row[0].value is None:
            continue
        if n >= len(values):
            sys.exit(f"{ws.title}: more data rows than values ({len(values)})")
        ws.cell(row=row[0].row, column=col, value=float(values[n]))
        n += 1
    if n != len(values):
        sys.exit(f"{ws.title}: {n} data rows, expected {len(values)}")


def document(doc, column: str, scenario: int, sources: list[str]) -> None:
    """Fill (or add) the row for this column in the Documentation sheet."""
    file_ = "RES/db_profiles_loader/update_etm_industry_profiles.py"
    what = ("sum of industry_*.input excl. steel and ICT" if column == "industry_other_e_demand"
            else sources[0])
    source = (f"ETM scenario {scenario} (electricity_profiles, retrieved "
              f"{datetime.now(timezone.utc):%Y-%m-%d}): {what}; weekdays aligned per "
              "calendar year (2024 +1 day, 2025 +3 days)")
    for row in doc.iter_rows(min_row=2):
        if row[0].value == column:
            row[3].value, row[4].value = file_, source
            return
    doc.append([column, COLUMNS[column][0], "Accumulative/interval-end", file_, source])


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("--scenario", type=int, default=DEFAULT_SCENARIO)
    p.add_argument("--dry-run", action="store_true", help="fetch and check, do not write")
    args = p.parse_args()

    curves = fetch_curves(args.scenario)
    profiles, problems = {}, []
    for column, (curve, sources) in curves.items():
        print(f"\nETM scenario {args.scenario}: {column} <- {len(sources)} column(s), "
              f"{curve.sum() / 1000:.1f} GWh")
        for c in sources:
            print(f"    {c}")
        reference = curve / curve.sum()
        for sheet, (year, days, quarter) in SHEETS.items():
            hourly = for_year(curve, year, days)
            values = to_quarter(hourly) if quarter else hourly
            profiles[(sheet, column)] = values
            problems += check(column, sheet, values, year,
                              reference if year == ETM_CALENDAR_YEAR and days == 365 else None)

    if problems:
        print("\nCHECK FAILED:\n  " + "\n  ".join(problems))
        sys.exit(1)
    print("\nChecks OK: annual sum 1, weekend lowest on Sat/Sun in every year, "
          "2023 identical in shape to the ETM curves.")
    if args.dry_run:
        return

    wb = openpyxl.load_workbook(WORKBOOK)
    for (sheet, column), values in profiles.items():
        write_column(wb[sheet], column, values)
    for column, (_, sources) in curves.items():
        document(wb["Documentation"], column, args.scenario, sources)
    wb.save(WORKBOOK)
    print(f"Written: {WORKBOOK}")


if __name__ == "__main__":
    main()

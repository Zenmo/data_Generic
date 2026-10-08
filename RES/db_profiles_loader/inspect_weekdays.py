"""(Re)build the check sheet 'Controle_weekdagen' in db_profiles.xlsx.

For each hourly profile sheet (profiles_2023_h, profiles_2024_h, profiles_2024_366_days_h,
profiles_2025_h) and each demand profile it writes two small tables with a line chart:

  1. Mean per weekday (Mon..Sun): the mean daily total on that weekday, relative to the
     mean daily total of the year. A weekday-driven profile should have its dip (or, for
     hot water, its peak) on Saturday/Sunday.
  2. Christmas period, 15 Dec - 31 Dec plus 1 - 7 Jan of the same sheet year: each day's
     total relative to the mean daily total of the year (same basis as table 1), so both
     the weekend rhythm and the holiday days are visible. The low-demand holiday week
     should fall in the last week of the year.

The quarter-hour sheets hold the same profiles (interpolated) and are not repeated.
Heat profiles also follow the weather, so their Christmas-period values mix holiday and
temperature effects.

Run after changing profiles (from this folder):
    python inspect_weekdays.py
Needs: numpy, openpyxl.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parent.parent / "db_profiles.xlsx"
SHEET = "Controle_weekdagen"

SOURCES = [  # (sheet, calendar year, days)
    ("profiles_2023_h", 2023, 365),
    ("profiles_2024_h", 2024, 365),
    ("profiles_2024_366_days_h", 2024, 366),
    ("profiles_2025_h", 2025, 365),
]
PROFILES = [
    "house_e_demand_other",
    "building_e_demand_other",
    "industry_other_e_demand",
    "ict_e_demand",
    "house_cooking_demand",
    "house_h_demand_hot_water",
    "building_h_demand",
    "industry_other_h_demand",
]
WEEKDAYS = ["ma", "di", "wo", "do", "vr", "za", "zo"]
BLOCK_ROWS = 48
WEEKEND_FILL = PatternFill("solid", fgColor="EDEDED")
HOLIDAY_FILL = PatternFill("solid", fgColor="FCE4D6")


def daily_totals(ws, days: int) -> dict[str, np.ndarray]:
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    cols = {p: header.index(p) for p in PROFILES if p in header}
    data = {p: [] for p in cols}
    n = 0
    for r in rows:
        if r[0] is None:
            continue
        for p, c in cols.items():
            v = r[c] if c < len(r) else None
            data[p].append(float(v) if isinstance(v, (int, float)) else 0.0)
        n += 1
        if n == days * 24:
            break
    return {p: np.array(v).reshape(days, 24).sum(axis=1) for p, v in data.items()}


def add_chart(ws, title, y_title, data_ref, cats_ref, anchor):
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.height, chart.width = 8.5, 22
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    # openpyxl marks axes as deleted by default; recent Excel versions then hide the
    # day names / dates on the x-axis and the values on the y-axis.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblSkip = 1
    chart.x_axis.title = "dag"
    ws.add_chart(chart, anchor)


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws["A1"] = "Controle weekdagen en kerstperiode (gegenereerd door RES/db_profiles_loader/inspect_weekdays.py)"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = ("Tabel 1: gemiddeld dagtotaal per weekdag / gemiddeld dagtotaal van het jaar. "
                "Tabel 2: dagtotaal / gemiddeld dagtotaal van het jaar (zelfde basis als tabel 1); "
                "januari = begin van hetzelfde jaar. "
                "Grijs = weekend, oranje = 25/26 dec en 1 jan. Kwartiertabbladen zijn gelijk (geïnterpoleerd).")

    for b, (sheet, year, days) in enumerate(SOURCES):
        totals = daily_totals(wb[sheet], days)
        profiles = [p for p in PROFILES if p in totals]
        dates = [date(year, 1, 1) + timedelta(days=i) for i in range(days)]
        top = 4 + b * BLOCK_ROWS

        ws.cell(row=top, column=1, value=f"{sheet} (kalenderjaar {year})").font = Font(bold=True, size=12)

        # Table 1: weekday means
        t1 = top + 1
        ws.cell(row=t1, column=1, value="weekdag").font = Font(bold=True)
        for j, p in enumerate(profiles, start=2):
            ws.cell(row=t1, column=j, value=p).font = Font(bold=True)
        for w in range(7):
            ws.cell(row=t1 + 1 + w, column=1, value=WEEKDAYS[w])
            for j, p in enumerate(profiles, start=2):
                d = totals[p]
                sel = [i for i in range(days) if dates[i].weekday() == w]
                ws.cell(row=t1 + 1 + w, column=j, value=round(float(d[sel].mean() / d.mean()), 4))
            if w >= 5:
                for j in range(1, len(profiles) + 2):
                    ws.cell(row=t1 + 1 + w, column=j).fill = WEEKEND_FILL
        add_chart(ws, f"{sheet}: gemiddelde per weekdag (1 = jaargemiddelde)", "relatief",
                  Reference(ws, min_col=2, max_col=1 + len(profiles), min_row=t1, max_row=t1 + 7),
                  Reference(ws, min_col=1, min_row=t1 + 1, max_row=t1 + 7),
                  f"K{t1}")

        # Table 2: Christmas period
        t2 = t1 + 10
        ws.cell(row=t2, column=1, value="dag").font = Font(bold=True)
        for j, p in enumerate(profiles, start=2):
            ws.cell(row=t2, column=j, value=p).font = Font(bold=True)
        window = ([i for i in range(days) if dates[i] >= date(year, 12, 15)]
                  + [i for i in range(days) if dates[i] <= date(year, 1, 7)])
        for k, i in enumerate(window):
            r = t2 + 1 + k
            ws.cell(row=r, column=1, value=f"{WEEKDAYS[dates[i].weekday()]} {dates[i].day}-{dates[i].month}")
            for j, p in enumerate(profiles, start=2):
                d = totals[p]
                ws.cell(row=r, column=j, value=round(float(d[i] / d.mean()), 4))
            fill = (HOLIDAY_FILL if (dates[i].month, dates[i].day) in ((12, 25), (12, 26), (1, 1))
                    else WEEKEND_FILL if dates[i].weekday() >= 5 else None)
            if fill:
                for j in range(1, len(profiles) + 2):
                    ws.cell(row=r, column=j).fill = fill
        add_chart(ws, f"{sheet}: 15 dec - 31 dec en 1 - 7 jan (1 = gemiddelde dag van het jaar)", "relatief",
                  Reference(ws, min_col=2, max_col=1 + len(profiles), min_row=t2, max_row=t2 + len(window)),
                  Reference(ws, min_col=1, min_row=t2 + 1, max_row=t2 + len(window)),
                  f"K{t2 + 8}")

    ws.column_dimensions["A"].width = 14
    for j in range(2, 2 + len(PROFILES)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 13
    wb.save(WORKBOOK)
    print(f"Written sheet {SHEET} in {WORKBOOK}")


if __name__ == "__main__":
    main()
